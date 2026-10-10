"""Export the catalog as a Word table, sorted by country.

Columns: Name & direct link to dataset | Country / Region | Powered by |
Catalyzed by | Financed by.

The "Powered by / Catalyzed by / Financed by" values are parsed out of each
project's `organizations` field. The name links to the project's page on the
live Fair Forward catalog (https://fair-forward.github.io/datasets/), which
resolves by the leading ui_N id, so it stays valid even for projects with
multiple dataset links or no public link.

Requires python-docx (`pip install python-docx`), which is not part of the
build requirements since this export runs on demand, not in the pipeline.
Root-level *.docx files are gitignored, so the default output stays local.

Usage:
    python scripts/export_catalog_table_docx.py \
        --input public/data/catalog.json \
        --output FAIR_Forward_Open_Data_by_Country.docx
"""

import argparse
import json
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

CATALOG_URL = "https://fair-forward.github.io/datasets/"

LABEL_RE = re.compile(
    r"(powered\s*by(?:\s*/\s*provided\s*by)?|provided\s*by|"
    r"cataly[sz]ed\s*by|financed\s*by|funded\s*by)\s*:?",
    re.IGNORECASE,
)
URL_IN_PARENS = re.compile(r"\s*\(\s*https?://[^)]*\)")


def clean_val(value):
    """Tidy a parsed org value: drop URL citations, collapse whitespace."""
    value = URL_IN_PARENS.sub("", value)
    value = re.sub(r"\s*\n\s*", " ", value)
    value = re.sub(r"\s{2,}", " ", value).strip()
    return value.lstrip(":").strip().strip(",").strip()


def parse_orgs(text):
    """Split an `organizations` blob into powered / catalyzed / financed.

    Handles labels on the same line or the next line, the
    "Powered by / Provided by" alias, and British/US spellings. Text with no
    recognisable label falls back to `powered`.
    """
    res = {"powered": "", "catalyzed": "", "financed": ""}
    if not text or not text.strip():
        return res
    matches = list(LABEL_RE.finditer(text))
    if not matches:
        res["powered"] = clean_val(text)
        return res
    for i, m in enumerate(matches):
        key = m.group(1).lower()
        if key.startswith(("powered", "provided")):
            slot = "powered"
        elif key.startswith("cataly"):
            slot = "catalyzed"
        else:
            slot = "financed"
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        val = clean_val(text[start:end])
        if val:
            res[slot] = val
    return res


def add_hyperlink(paragraph, url, text, size_pt=9.5):
    """Append a clickable hyperlink run to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)

    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(rfonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size_pt * 2)))
    rpr.append(sz)

    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_cell_shading(cell, fill_hex):
    """Apply a solid background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def style_cell_text(cell, text, bold=False, size_pt=9.5, color=None):
    """Replace a cell's text with a single styled run."""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color


def build_rows(projects):
    """Return display rows sorted by first country, then title."""
    rows = []
    for p in projects:
        countries = [c.strip() for c in p.get("countries", []) if c.strip()]
        orgs = parse_orgs(p.get("organizations", ""))
        slug = p.get("slug") or p.get("id")
        rows.append(
            {
                "title": p.get("title", "").strip(),
                "url": f"{CATALOG_URL}?project={slug}",
                "country": ", ".join(countries),
                "sort_country": countries[0] if countries else "ZZZ",
                "powered": orgs["powered"],
                "catalyzed": orgs["catalyzed"],
                "financed": orgs["financed"],
            }
        )
    rows.sort(key=lambda r: (r["sort_country"].lower(), r["title"].lower()))
    return rows


COLUMNS = [
    ("Name & direct link to dataset", Cm(5.6)),
    ("Country / Region", Cm(2.6)),
    ("Powered by", Cm(3.5)),
    ("Catalyzed by", Cm(3.1)),
    ("Financed by", Cm(2.2)),
]
HEADER_FILL = "D9E2F3"
BODY_SIZE = 9.5


def build_document(rows):
    doc = Document()

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    heading = doc.add_paragraph()
    run = heading.add_run("FAIR Forward - Catalog of Open Data & Use Cases")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)

    sub = doc.add_paragraph()
    sub_run = sub.add_run(
        f"Open datasets and AI use cases, sorted by country ({len(rows)} entries)."
    )
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    # Fixed layout so the column widths are honoured by Word.
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    header_cells = table.rows[0].cells
    for idx, (label, width) in enumerate(COLUMNS):
        cell = header_cells[idx]
        cell.width = width
        set_cell_shading(cell, HEADER_FILL)
        style_cell_text(cell, label, bold=True, size_pt=10)

    for row in rows:
        cells = table.add_row().cells
        for idx, (_, width) in enumerate(COLUMNS):
            cells[idx].width = width

        # Name + hyperlink
        name_cell = cells[0]
        name_cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        add_hyperlink(name_cell.paragraphs[0], row["url"], row["title"], BODY_SIZE)

        style_cell_text(cells[1], row["country"], size_pt=BODY_SIZE)
        style_cell_text(cells[2], row["powered"], size_pt=BODY_SIZE)
        style_cell_text(cells[3], row["catalyzed"], size_pt=BODY_SIZE)
        style_cell_text(cells[4], row["financed"], size_pt=BODY_SIZE)

    return doc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="public/data/catalog.json")
    parser.add_argument("--output", default="FAIR_Forward_Open_Data_by_Country.docx")
    args = parser.parse_args()

    with open(args.input, encoding="utf-8") as f:
        catalog = json.load(f)

    rows = build_rows(catalog["projects"])
    doc = build_document(rows)
    doc.save(args.output)

    print(f"Wrote {args.output} with {len(rows)} rows.")
    blanks = [
        r for r in rows
        if not (r["powered"] and r["catalyzed"] and r["financed"])
    ]
    if blanks:
        print(f"\n{len(blanks)} row(s) missing at least one org field:")
        for r in blanks:
            missing = [k for k in ("powered", "catalyzed", "financed") if not r[k]]
            print(f"  - {r['title'][:55]!r}: missing {', '.join(missing)}")


if __name__ == "__main__":
    main()
